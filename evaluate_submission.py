from pathlib import Path
import json
import sys
import tempfile
import shutil
import zipfile

import pandas as pd
import numpy as np
from scipy.stats import chi2
from openpyxl import load_workbook
from statsmodels.stats.outliers_influence import variance_inflation_factor


# ---------------------------------------------------------
# Project paths
# ---------------------------------------------------------

PROJECT_ROOT = Path(__file__).resolve().parent

DATASET_DIR = PROJECT_ROOT / "datasets"
PROFILE_DIR = PROJECT_ROOT / "dataset_profiles"


# ---------------------------------------------------------
# Helper functions
# ---------------------------------------------------------

def load_dataset(dataset_id):

    if dataset_id in ("01", "02"):

        dataset_file = (
            DATASET_DIR
            / f"ML_Dataset_{dataset_id}.csv"
        )

    else:

        dataset_file = (
            DATASET_DIR
            / "variants"
            / f"ML_Dataset_{dataset_id}.csv"
        )

    if not dataset_file.exists():
        raise FileNotFoundError(
            f"Dataset not found: {dataset_file}"
        )

    return pd.read_csv(dataset_file)


def load_profile(dataset_id):

    if dataset_id in ("01", "02"):

        profile_file = (
            PROFILE_DIR
            / f"dataset_{dataset_id.lower()}.json"
        )

    else:

        profile_file = (
            PROFILE_DIR
            / "variants"
            / f"ML_Dataset_{dataset_id}.json"
        )

    if not profile_file.exists():
        raise FileNotFoundError(
            f"Profile not found: {profile_file}"
        )

    with open(
        profile_file,
        "r",
        encoding="utf-8"
    ) as f:

        return json.load(f)


def calculate_highest_correlation(df, predictors):
    X = df[predictors]

    corr = X.corr().abs()

    best_value = -1
    best_var1 = None
    best_var2 = None

    for i in range(len(predictors)):
        for j in range(i + 1, len(predictors)):

            var1 = predictors[i]
            var2 = predictors[j]

            value = float(corr.loc[var1, var2])

            if value > best_value:
                best_value = value
                best_var1 = var1
                best_var2 = var2

    return best_var1, best_var2, best_value


def calculate_vif(df, predictors):

    X = df[predictors].copy()
    X.insert(0, "const", 1.0)

    results = []

    for i, column in enumerate(X.columns):

        if column == "const":
            continue

        vif = variance_inflation_factor(
            X.values,
            i
        )

        results.append(
            {
                "variable": column,
                "vif": float(vif)
            }
        )

    results.sort(
        key=lambda x: x["vif"],
        reverse=True
    )

    return results


def compare_number(
    reported,
    actual,
    tolerance=0
):
    try:
        reported = float(reported)
        actual = float(actual)

        difference = abs(
            reported - actual
        )

        return {
            "reported": reported,
            "actual": actual,
            "difference": difference,
            "verified": difference <= tolerance
        }

    except (TypeError, ValueError):

        return {
            "reported": reported,
            "actual": actual,
            "difference": None,
            "verified": False
        }


def normalize_text(value):
    if value is None:
        return ""

    return str(value).strip().lower()
def evaluate_dimensionality_reduction(ws, df, profile):
    """
    Independently evaluate objective dimensionality-reduction results.
    """

    results = []

    design = profile.get(
        "dimensionality_reduction"
    )

    if not design:
        return []

    reference = profile.get(
        "dimensionality_reduction_reference",
        {}
    )

    rules = profile.get(
        "dimensionality_reduction_evaluation",
        {}
    )

    expected_dimensions = design[
        "original_dimensions"
    ]

    # ---------------------------------------------------------
    # Read student entries
    # ---------------------------------------------------------

    reported_used = ws["B4"].value
    reported_method = ws["B5"].value
    reported_original = ws["B6"].value
    reported_final = ws["B7"].value
    reported_variance = ws["B8"].value

    method = normalize_text(reported_method)

    # ---------------------------------------------------------
    # Detect untouched template placeholders
    # ---------------------------------------------------------

    placeholder_values = {
        "yes / no",
        "e.g., pca / factor analysis / other / not applicable",
        "e.g., pca / other / not applicable",
        "enter value",
        "enter value if applicable"
    }

    normalized_used = normalize_text(
        reported_used
    )

    normalized_method = normalize_text(
        reported_method
    )

    template_not_completed = (
        normalized_used in placeholder_values
        or normalized_method in placeholder_values
    )

    if template_not_completed:

        return [
            {
                "check": "Dimensionality-reduction section",
                "reported": reported_method,
                "actual": "Faculty review",
                "verified": None
            }
        ]

    # Original number of dimensions
    # ---------------------------------------------------------

    original_check = compare_number(
        reported_original,
        expected_dimensions,
        tolerance=0
    )

    results.append(
        {
            "check": "Original number of dimensions",
            "reported": reported_original,
            "actual": expected_dimensions,
            "verified": original_check["verified"]
        }
    )

    # ---------------------------------------------------------
    # PCA
    # ---------------------------------------------------------

    if "pca" in method:

        cumulative = reference["pca"][
            "cumulative_variance_percent"
        ]

        pca_variance_verified = False
        actual_variance = None

        try:
            component_count = int(
                float(reported_final)
            )

            if 1 <= component_count <= len(cumulative):

                actual_variance = cumulative[
                    component_count - 1
                ]

                comparison = compare_number(
                    reported_variance,
                    actual_variance,
                    tolerance=rules[
                        "pca_variance_tolerance"
                    ]
                )

                pca_variance_verified = (
                    comparison["verified"]
                )

        except (TypeError, ValueError):
            pass

        results.append(
            {
                "check": "PCA variance explained",
                "reported": reported_variance,
                "actual": actual_variance,
                "component_count": reported_final,
                "verified": pca_variance_verified
            }
        )

    # ---------------------------------------------------------
        # ---------------------------------------------------------
    # Factor Analysis
    # ---------------------------------------------------------

    elif (
        "factor" in method
        or method == "fa"
    ):

        fa_reference = reference[
            "factor_analysis"
        ]

        # Student-reported FA suitability statistics
        reported_kmo = ws["B11"].value
        reported_bartlett_chi = ws["B12"].value
        reported_bartlett_df = ws["B13"].value
        reported_bartlett_p = ws["B14"].value

        # -----------------------------------------------------
        # KMO
        # -----------------------------------------------------

        kmo_check = compare_number(
            reported_kmo,
            fa_reference["kmo"],
            tolerance=rules["kmo_tolerance"]
        )

        results.append(
            {
                "check": "Factor Analysis KMO",
                "reported": reported_kmo,
                "actual": fa_reference["kmo"],
                "verified": kmo_check["verified"]
            }
        )

        # -----------------------------------------------------
        # Bartlett's test — chi-square
        # -----------------------------------------------------

        bartlett_chi_check = compare_number(
            reported_bartlett_chi,
            fa_reference["bartlett_chi_square"],
            tolerance=rules[
                "bartlett_chi_square_tolerance"
            ]
        )

        results.append(
            {
                "check": "Bartlett's test chi-square",
                "reported": reported_bartlett_chi,
                "actual": fa_reference[
                    "bartlett_chi_square"
                ],
                "verified": bartlett_chi_check["verified"]
            }
        )

        # -----------------------------------------------------
        # Bartlett's test — degrees of freedom
        # -----------------------------------------------------

        bartlett_df_check = compare_number(
            reported_bartlett_df,
            fa_reference["bartlett_df"],
            tolerance=0
        )

        results.append(
            {
                "check": "Bartlett's test df",
                "reported": reported_bartlett_df,
                "actual": fa_reference["bartlett_df"],
                "verified": bartlett_df_check["verified"]
            }
        )

        # -----------------------------------------------------
        # Bartlett's test — p-value
        # -----------------------------------------------------
        #
        # The reference is expressed as p < .001.
        # We therefore verify that the student's reported
        # p-value is below .001 rather than requiring an
        # exact numeric match.
        # -----------------------------------------------------

        p_verified = False

        try:
            p_text = str(
                reported_bartlett_p
            ).strip().lower()

            reference_p = fa_reference[
                "bartlett_p_value_less_than"
            ]

            if p_text.startswith("<"):
                # Example: "< .001" or "< 0.001"
                reported_p = float(
                    p_text[1:].strip()
                )

                p_verified = (
                    reported_p <= reference_p
                )

            else:
                # Example: 0.0005 or 0.0009
                reported_p = float(
                    p_text.replace("=", "").strip()
                )

                p_verified = (
                    reported_p < reference_p
                )

        except (TypeError, ValueError):
            pass

        results.append(
            {
                "check": "Bartlett's test p-value",
                "reported": reported_bartlett_p,
                "actual": (
                    f"< {fa_reference['bartlett_p_value_less_than']}"
                ),
                "verified": p_verified
            }
        )
    # ---------------------------------------------------------
    # Other / Not applicable
    # ---------------------------------------------------------

    else:

        results.append(
            {
                "check": "Dimensionality-reduction method",
                "reported": reported_method,
                "actual": "Faculty review",
                "verified": None
            }
        )

    return results
# ---------------------------------------------------------
# Model Evaluation
# ---------------------------------------------------------

def evaluate_model_evaluation(ws, profile):
    """
    Independently evaluate objective model-evaluation evidence.

    A canonical OLS baseline is available for Dataset 01.
    Alternative defensible models remain faculty-review items.
    """

    results = []

    reference = profile.get("regression_reference")
    rules = profile.get("regression_evaluation", {})

    if not reference or not rules.get(
        "verify_baseline_metrics",
        False
    ):
        return results

    # -----------------------------------------------------
    # Read final-selection fields
    # -----------------------------------------------------

    selected_approach = normalize_text(
        ws["B18"].value
    )

    primary_metric = normalize_text(
        ws["B19"].value
    )

    final_score = ws["B20"].value

    # -----------------------------------------------------
    # Detect canonical OLS baseline
    # -----------------------------------------------------

    is_ols_baseline = (
        "ols" in selected_approach
        or "ordinary least squares" in selected_approach
        or "linear regression" in selected_approach
    )

    # -----------------------------------------------------
    # Alternative model:
    # faculty review rather than automatic judgement.
    # -----------------------------------------------------

    if not is_ols_baseline:

        results.append(
            {
                "check": "Model selection / evaluation",
                "reported": ws["B18"].value,
                "actual": "Faculty review",
                "verified": None
            }
        )

        return results

    # -----------------------------------------------------
    # Determine primary metric
    # -----------------------------------------------------

    if "adjusted" in primary_metric:

        expected = reference[
            "adjusted_r_squared"
        ]

        tolerance = rules[
            "adjusted_r_squared_tolerance"
        ]

        check_name = "OLS adjusted R-squared"

    elif (
        "r2" in primary_metric
        or "r-squared" in primary_metric
        or "r squared" in primary_metric
        or "r²" in primary_metric
    ):

        expected = reference[
            "r_squared"
        ]

        tolerance = rules[
            "r_squared_tolerance"
        ]

        check_name = "OLS R-squared"

    elif "rmse" in primary_metric:

        expected = reference["rmse"]

        tolerance = rules[
            "rmse_tolerance"
        ]

        check_name = "OLS RMSE"

    else:

        results.append(
            {
                "check": "OLS primary evaluation metric",
                "reported": ws["B19"].value,
                "actual": "Faculty review",
                "verified": None
            }
        )

        return results

    # -----------------------------------------------------
    # Verify reported score
    # -----------------------------------------------------

    comparison = compare_number(
        final_score,
        expected,
        tolerance=tolerance
    )

    results.append(
        {
            "check": check_name,
            "reported": final_score,
            "actual": expected,
            "verified": comparison["verified"]
        }
    )

    return results
# ---------------------------------------------------------
# Classification Evaluation
# ---------------------------------------------------------

def evaluate_classification(ws, profile):
    """
    Independently evaluate the canonical classification baseline.

    Logistic Regression + ROC-AUC is the reference baseline for
    Dataset 02. Alternative defensible models remain faculty-review
    items.
    """

    results = []

    reference = profile.get(
        "classification_reference"
    )

    rules = profile.get(
        "classification_evaluation",
        {}
    )

    if not reference or not rules.get(
        "verify_baseline_metric",
        False
    ):
        return results

    # -----------------------------------------------------
    # Read final-selection fields
    # -----------------------------------------------------

    selected_approach = normalize_text(
        ws["B18"].value
    )

    primary_metric = normalize_text(
        ws["B19"].value
    )

    final_score = ws["B20"].value

    # -----------------------------------------------------
    # Detect Logistic Regression baseline
    # -----------------------------------------------------

    is_logistic_baseline = (
        "logistic" in selected_approach
    )

    # -----------------------------------------------------
    # Alternative model → faculty review
    # -----------------------------------------------------

    if not is_logistic_baseline:

        results.append(
            {
                "check": "Classification model selection",
                "reported": ws["B18"].value,
                "actual": "Faculty review",
                "verified": None
            }
        )

        return results

    # -----------------------------------------------------
    # Verify ROC-AUC
    # -----------------------------------------------------

    if (
        "roc" not in primary_metric
        or "auc" not in primary_metric
    ):

        results.append(
            {
                "check": "Classification evaluation metric",
                "reported": ws["B19"].value,
                "actual": "Faculty review",
                "verified": None
            }
        )

        return results

    expected = reference["roc_auc"]

    comparison = compare_number(
        final_score,
        expected,
        tolerance=rules[
            "roc_auc_tolerance"
        ]
    )

    results.append(
        {
            "check": "Logistic Regression ROC-AUC",
            "reported": final_score,
            "actual": expected,
            "verified": comparison["verified"]
        }
    )

    return results

    # ---------------------------------------------------------
# Clustering Evaluation
# ---------------------------------------------------------

def evaluate_clustering(ws, profile):

    results = []

    reference = profile.get(
        "clustering_reference"
    )

    rules = profile.get(
        "clustering_evaluation",
        {}
    )

    if not reference or not rules.get(
        "verify_validation_score",
        False
    ):
        return results

    # -----------------------------------------------------
    # Read reported clustering fields
    # -----------------------------------------------------

    performed = ws["B4"].value
    method = normalize_text(
        ws["B5"].value
    )
    clusters = ws["B7"].value
    metric = normalize_text(
        ws["B9"].value
    )
    score = ws["B10"].value

    # -----------------------------------------------------
    # If clustering was not performed
    # -----------------------------------------------------

    if normalize_text(performed) in [
        "no",
        "n"
    ]:

        results.append(
            {
                "check": "Clustering performed",
                "reported": performed,
                "actual": "Faculty review",
                "verified": None
            }
        )

        return results

    # -----------------------------------------------------
    # Only silhouette score can be independently verified
    # -----------------------------------------------------

    if (
        "silhouette" not in metric
        or clusters is None
        or score is None
    ):

        results.append(
            {
                "check": "Clustering validation",
                "reported": score,
                "actual": "Faculty review",
                "verified": None
            }
        )

        return results

    try:
        k = int(clusters)
    except (TypeError, ValueError):

        results.append(
            {
                "check": "Clustering validation",
                "reported": score,
                "actual": "Faculty review",
                "verified": None
            }
        )

        return results

    # -----------------------------------------------------
    # Identify reference method
    # -----------------------------------------------------

    if "kmeans" in method or "k-means" in method:

        reference_scores = reference[
            "kmeans_reference"
        ][
            "silhouette_scores"
        ]

        expected = reference_scores.get(
            str(k)
        )

    elif (
        "hierarchical" in method
        or "ward" in method
    ):

        reference_scores = reference[
            "hierarchical_ward_reference"
        ][
            "silhouette_scores"
        ]

        expected = reference_scores.get(
            str(k)
        )

    else:

        results.append(
            {
                "check": "Clustering method",
                "reported": ws["B5"].value,
                "actual": "Faculty review",
                "verified": None
            }
        )

        return results

    # -----------------------------------------------------
    # Method / K combination not in reference
    # -----------------------------------------------------

    if expected is None:

        results.append(
            {
                "check": "Clustering validation",
                "reported": score,
                "actual": "Faculty review",
                "verified": None
            }
        )

        return results

    # -----------------------------------------------------
    # Verify silhouette score
    # -----------------------------------------------------

    comparison = compare_number(
        score,
        expected,
        tolerance=rules[
            "silhouette_tolerance"
        ]
    )

    results.append(
        {
            "check": "Clustering silhouette score",
            "reported": score,
            "actual": expected,
            "verified": comparison[
                "verified"
            ]
        }
    )

    return results
# ---------------------------------------------------------
# Main evaluation
# ---------------------------------------------------------

def evaluate_submission(zip_path, dataset_id="01", return_results=False):

    zip_path = Path(zip_path)

    print()
    print("=" * 70)
    print("ML ASSIGNMENT — INDEPENDENT EVALUATION")
    print("=" * 70)

    if not zip_path.exists():
        print("ERROR: Submission ZIP not found.")
        print(zip_path)
        return False

    # -----------------------------------------------------
    # Load dataset and profile
    # -----------------------------------------------------

    df = load_dataset(dataset_id)
    profile = load_profile(dataset_id)

    predictors = profile[
        "analysis_design"
    ][
        "diagnostic_predictors"
    ]

    print()
    print(f"Dataset: {profile['dataset_file']}")
    print(f"Rows: {df.shape[0]}")
    print(f"Columns: {df.shape[1]}")

    # -----------------------------------------------------
    # Extract submission
    # -----------------------------------------------------

    temp_dir = Path(
        tempfile.mkdtemp(
            prefix="ml_evaluation_"
        )
    )

    try:

        with zipfile.ZipFile(
            zip_path,
            "r"
        ) as z:

            z.extractall(temp_dir)

        excel_files = list(
            temp_dir.rglob(
                "Results_Template.xlsx"
            )
        )

        if len(excel_files) != 1:

            print()
            print(
                "ERROR: Expected exactly one Results_Template.xlsx"
            )

            return False

        excel_file = excel_files[0]

        print(
            f"Results file: {excel_file.name}"
        )

        # -------------------------------------------------
        # Read workbook
        # -------------------------------------------------

        wb = load_workbook(
            excel_file,
            data_only=True
        )

        # -------------------------------------------------
        # Results
        # -------------------------------------------------

        results = []
        # =================================================
        # 1A. Dimensionality Reduction
        # =================================================

        ws = wb["Dimensionality_Reduction"]

        dimensionality_results = (
            evaluate_dimensionality_reduction(
                ws,
                df,
                profile
            )
        )

        results.extend(dimensionality_results)

                # =================================================
        # 1B. Model Evaluation
        # =================================================

        ws = wb["Model_Evaluation"]

        model_evaluation_results = (
            evaluate_model_evaluation(
                ws,
                profile
            )
        )

        results.extend(model_evaluation_results)
                # =================================================
        # 1C. Classification
        # =================================================

        ws = wb["Model_Evaluation"]

        classification_results = (
            evaluate_classification(
                ws,
                profile
            )
        )

        results.extend(classification_results)

                # =================================================
        # 1D. Clustering
        # =================================================

        ws = wb["Clustering"]

        clustering_results = (
            evaluate_clustering(
                ws,
                profile
            )
        )

        results.extend(clustering_results)

        # =================================================
        # 1. Dataset dimensions
        # =================================================

        ws = wb["EDA_Results"]

        reported_rows = ws["B5"].value
        reported_columns = ws["B6"].value
        reported_missing = ws["B7"].value
        reported_duplicates = ws["B8"].value

        checks = [
            (
                "Number of observations",
                reported_rows,
                df.shape[0]
            ),
            (
                "Number of variables",
                reported_columns,
                df.shape[1]
            ),
            (
                "Missing values",
                reported_missing,
                int(df.isna().sum().sum())
            ),
            (
                "Duplicate observations",
                reported_duplicates,
                int(df.duplicated().sum())
            ),
        ]

        for name, reported, actual in checks:

            comparison = compare_number(
                reported,
                actual
            )

            results.append(
                {
                    "check": name,
                    **comparison
                }
            )

        # =================================================
        # 2. Highest correlation
        # =================================================

        ws = wb["Multicollinearity"]

        reported_corr = ws["B5"].value
        reported_var1 = ws["C5"].value
        reported_var2 = ws["D5"].value

        actual_var1, actual_var2, actual_corr = (
            calculate_highest_correlation(
                df,
                predictors
            )
        )

        correlation_tolerance = profile[
            "evaluation_rules"
        ][
            "correlation_tolerance"
        ]

        value_check = compare_number(
            reported_corr,
            actual_corr,
            correlation_tolerance
        )

        reported_pair = {
            normalize_text(reported_var1),
            normalize_text(reported_var2)
        }

        actual_pair = {
            normalize_text(actual_var1),
            normalize_text(actual_var2)
        }

        pair_match = (
            reported_pair == actual_pair
        )

        correlation_verified = (
            value_check["verified"]
            and pair_match
        )

        results.append(
            {
                "check": "Highest predictor correlation",
                "reported": reported_corr,
                "actual": actual_corr,
                "difference": value_check[
                    "difference"
                ],
                "reported_variable_1": reported_var1,
                "reported_variable_2": reported_var2,
                "actual_variable_1": actual_var1,
                "actual_variable_2": actual_var2,
                "value_verified": value_check[
                    "verified"
                ],
                "variables_verified": pair_match,
                "verified": correlation_verified
            }
        )

        # =================================================
        # 3. Highest VIF
        # =================================================

        reported_vif = ws["B6"].value
        reported_vif_variable = ws["C6"].value

        vif_results = calculate_vif(
            df,
            predictors
        )

        actual_vif_variable = vif_results[0][
            "variable"
        ]

        actual_vif = vif_results[0][
            "vif"
        ]

        vif_tolerance = profile[
            "evaluation_rules"
        ][
            "vif_tolerance"
        ]

        vif_value_check = compare_number(
            reported_vif,
            actual_vif,
            vif_tolerance
        )

        vif_variable_match = (
            normalize_text(
                reported_vif_variable
            )
            ==
            normalize_text(
                actual_vif_variable
            )
        )

        vif_verified = (
            vif_value_check["verified"]
            and vif_variable_match
        )

        results.append(
            {
                "check": "Highest predictor VIF",
                "reported": reported_vif,
                "actual": actual_vif,
                "difference": vif_value_check[
                    "difference"
                ],
                "reported_variable": reported_vif_variable,
                "actual_variable": actual_vif_variable,
                "value_verified": vif_value_check[
                    "verified"
                ],
                "variable_verified": vif_variable_match,
                "verified": vif_verified
            }
        )
                # =================================================
        # 4. Number of variables affected by multicollinearity
        # =================================================

        reported_affected = ws["B7"].value

        vif_threshold = profile[
            "hidden_characteristics"
        ][
            "vif_threshold_flag"
        ]

        actual_affected = sum(
            item["vif"] >= vif_threshold
            for item in vif_results
        )

        affected_check = compare_number(
            reported_affected,
            actual_affected
        )

        results.append(
            {
                "check": "Variables affected by multicollinearity",
                "reported": reported_affected,
                "actual": actual_affected,
                "difference": affected_check[
                    "difference"
                ],
                "criterion": f"VIF >= {vif_threshold}",
                "verified": affected_check[
                    "verified"
                ]
            }
        )

        # -------------------------------------------------
        # Print results
        # -------------------------------------------------

        print()
        print("-" * 70)
        print("VERIFICATION RESULTS")
        print("-" * 70)

        for result in results:

            if result["verified"] is True:
                status = "✓ VERIFIED"
            elif result["verified"] is False:
                status = "✗ REVIEW"
            else:
                status = "⚠ FACULTY REVIEW"

            print(
                f"{result['check']:<35} {status}"
            )

            print(
                f"  Reported: {result['reported']}"
            )

            print(
                f"  Actual:   {result['actual']}"
            )

        print()
        print("=" * 70)

        verified_count = sum(
            r["verified"] is True
            for r in results
        )

        print(
            f"Checks verified: "
            f"{verified_count}/{len(results)}"
        )

        print("=" * 70)
        print()

        if return_results:
            return results

        if return_results:
            return results

        return True

    finally:

        shutil.rmtree(
            temp_dir,
            ignore_errors=True
        )


# ---------------------------------------------------------
# Command-line interface
# ---------------------------------------------------------

if __name__ == "__main__":

    if len(sys.argv) not in (2, 3):

        print()
        print(
            "Usage:"
        )
        print(
            "python evaluate_submission.py <submission.zip> [dataset_id]"
        )
        print()

        sys.exit(1)

    submission_zip = sys.argv[1]

    dataset_id = (
        sys.argv[2]
        if len(sys.argv) == 3
        else "01"
    )

    success = evaluate_submission(
        submission_zip,
        dataset_id
    )

    sys.exit(
        0 if success else 1
    )